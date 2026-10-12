"""Excel → PDF through LibreOffice, with optional sheet selection and page-setup overrides."""
from __future__ import annotations

import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Sequence

import openpyxl

from .common import InvalidInputError, JobContext, Result
from .office_engine import libreoffice_to_pdf, validate_office_file

SUPPORTED = (".xlsx", ".xls", ".ods")  # .xlsm (macros) is intentionally not accepted
LARGE_CELLS = 1_000_000


@dataclass
class SheetInfo:
    name: str
    hidden: bool
    rows: int
    cols: int


def list_sheets(path: Path) -> list[SheetInfo]:
    """Sheet names for ``.xlsx`` files (the GUI uses this for the selection list)."""
    path = Path(path)
    validate_office_file(path, (".xlsx",))
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception as exc:
        raise InvalidInputError(f"'{path.name}' could not be read as an Excel workbook.") from exc
    try:
        return [SheetInfo(ws.title, ws.sheet_state != "visible", ws.max_row or 0, ws.max_column or 0)
                for ws in wb.worksheets]
    finally:
        wb.close()


def excel_to_pdf(src: Path, output: Path, *, sheets: Optional[Sequence[str]] = None,
                 orientation: str = "keep", fit_to_width: bool = False, overwrite: bool = False,
                 ctx: JobContext | None = None) -> Result:
    """``sheets`` – names to export (``None`` = all). ``orientation`` – ``keep|portrait|landscape``.

    With no overrides the original file is converted untouched and the workbook's own print areas, margins and
    scaling are respected.  Overrides require re-saving the workbook with openpyxl, which can drop charts,
    pivot tables and some formatting – a warning is added in that case.
    """
    ctx = ctx or JobContext()
    src = Path(src)
    validate_office_file(src, SUPPORTED)
    warnings: list[str] = []
    needs_edit = bool(sheets) or orientation != "keep" or fit_to_width
    if needs_edit and src.suffix.lower() != ".xlsx":
        raise InvalidInputError("Sheet selection and page-setup overrides are only available for .xlsx files.")
    if src.suffix.lower() == ".xlsx":
        infos = list_sheets(src)
        if sheets:
            unknown = [s for s in sheets if s not in {i.name for i in infos}]
            if unknown:
                raise InvalidInputError(f"Sheet not found: {', '.join(unknown)}")
        hidden = [i.name for i in infos if i.hidden]
        if hidden:
            warnings.append(f"Hidden sheet(s) are not exported: {', '.join(hidden)}.")
        big = [i.name for i in infos if i.rows * i.cols > LARGE_CELLS]
        if big:
            warnings.append(f"Large sheet(s) ({', '.join(big)}) may produce very many pages or take a long time.")
    if not needs_edit:
        res = libreoffice_to_pdf(src, output, overwrite=overwrite, ctx=ctx)
    else:
        with tempfile.TemporaryDirectory(prefix="pdftoolkit_xl_") as td:
            edited = Path(td) / "edited.xlsx"
            wb = openpyxl.load_workbook(src)  # formulas kept; LibreOffice recalculates on load
            try:
                chosen = set(sheets) if sheets else {ws.title for ws in wb.worksheets}
                first_visible = next((ws for ws in wb.worksheets if ws.title in chosen), None)
                for ws in wb.worksheets:
                    if ws.title not in chosen:
                        ws.sheet_state = "hidden"  # hide instead of delete so formulas stay valid
                    else:
                        if orientation in ("portrait", "landscape"):
                            ws.page_setup.orientation = orientation
                        if fit_to_width:
                            ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
                            ws.page_setup.fitToWidth = 1
                            ws.page_setup.fitToHeight = 0
                if first_visible is not None:
                    first_visible.sheet_state = "visible"
                    wb.active = wb.worksheets.index(first_visible)
                wb.save(edited)
            finally:
                wb.close()
            res = libreoffice_to_pdf(edited, output, overwrite=overwrite, ctx=ctx)
        warnings.append("The workbook was re-saved before conversion because you chose sheet/page-setup options; "
                        "charts, pivot tables and some formatting may be missing from the PDF.")
    res.warnings = warnings + res.warnings
    res.warnings.append("Page breaks follow the workbook's print settings; very wide sheets split across pages "
                        "unless 'fit to one page wide' is selected.")
    return res
