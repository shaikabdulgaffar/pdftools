"""Conversion tests.  Office tests need LibreOffice and are skipped (not faked) when it is absent."""
import pymupdf
import pytest
from PIL import Image

from core.common import (EngineMissingError, InvalidInputError, JobContext, OperationCancelled, ToolkitError)
from core.excel_to_pdf import excel_to_pdf, list_sheets
from core.html_to_pdf import html_to_pdf, scan_html
from core.image_to_pdf import LayoutOptions, images_to_pdf
from core.office_engine import find_soffice, require_soffice
from core.pdf_image_pdf import pdf_to_image_to_pdf
from core.pdf_to_image import pdf_to_images
from core.powerpoint_to_pdf import powerpoint_to_pdf
from core.word_to_pdf import word_to_pdf
from conftest import make_pdf, page_texts

needs_lo = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed (manual check required)")


# ------------------------------------------------------------------ image -> PDF
def test_images_to_pdf_formats_and_pages(tmp_path, img_dir):
    files = [img_dir / n for n in ("photo.jpg", "alpha.png", "palette.png", "pic.webp", "pic.bmp", "multi.tif")]
    out = tmp_path / "i.pdf"
    res = images_to_pdf(files, out)
    assert res.details["pages"] == 7  # multipage TIFF contributes 2 pages
    with pymupdf.open(out) as d:
        assert all(abs(p.rect.width - 595) < 1 or abs(p.rect.width - 842) < 1 for p in d)


def test_jpeg_is_not_reencoded(tmp_path, img_dir):
    out = tmp_path / "j.pdf"
    images_to_pdf([img_dir / "photo.jpg"], out)
    with pymupdf.open(out) as d:
        info = d.extract_image(d.get_page_images(0)[0][0])
    assert info["ext"] in ("jpeg", "jpg")
    assert info["image"] == (img_dir / "photo.jpg").read_bytes()  # byte-identical stream


def test_image_layout_options(tmp_path, img_dir):
    src = img_dir / "photo.jpg"  # 800x600, landscape
    out = tmp_path / "l.pdf"
    images_to_pdf([src], out, LayoutOptions(page_size="letter", orientation="portrait", margin_mm=20))
    with pymupdf.open(out) as d:
        assert (round(d[0].rect.width), round(d[0].rect.height)) == (612, 792)
        r = d[0].get_image_rects(d[0].get_images()[0][0])[0]
        m = 20 * 72 / 25.4
        assert r.x0 >= m - 1 and r.x1 <= 612 - m + 1 and abs(r.width / r.height - 800 / 600) < 0.01
    out2 = tmp_path / "l2.pdf"
    images_to_pdf([src], out2, LayoutOptions(page_size="image", margin_mm=0, placement="actual"))
    with pymupdf.open(out2) as d:  # page equals image size at 96 dpi fallback
        assert abs(d[0].rect.width - 800 * 72 / 96) < 2
    out3 = tmp_path / "l3.pdf"
    images_to_pdf([src], out3, LayoutOptions(page_size="a4", orientation="landscape"))
    with pymupdf.open(out3) as d:
        assert d[0].rect.width > d[0].rect.height


def test_images_to_pdf_errors(tmp_path, img_dir):
    with pytest.raises(InvalidInputError):
        images_to_pdf([], tmp_path / "x.pdf")
    with pytest.raises(InvalidInputError, match="not a readable image"):
        images_to_pdf([img_dir / "broken.png"], tmp_path / "x.pdf")
    txt = tmp_path / "a.txt"
    txt.write_text("x")
    with pytest.raises(InvalidInputError, match="supported"):
        images_to_pdf([txt], tmp_path / "x.pdf")
    with pytest.raises(InvalidInputError, match="margins"):
        images_to_pdf([img_dir / "photo.jpg"], tmp_path / "x.pdf", LayoutOptions(margin_mm=200))
    assert not (tmp_path / "x.pdf").exists()


# ------------------------------------------------------------------ PDF -> image
def test_pdf_to_images_png_jpeg_ranges(tmp_path, pdf_factory):
    src = pdf_factory("doc.pdf", pages=4)
    res = pdf_to_images(src, tmp_path / "png", fmt="png", dpi=72)
    assert [p.name for p in res.outputs] == [f"doc_page_00{i}.png" for i in range(1, 5)]
    with Image.open(res.outputs[0]) as im:
        assert im.size == (595, 842) and im.format == "PNG"
    res = pdf_to_images(src, tmp_path / "jpg", fmt="jpeg", dpi=144, pages_spec="2-3", jpeg_quality=60)
    assert [p.name for p in res.outputs] == ["doc_page_002.jpg", "doc_page_003.jpg"]
    with Image.open(res.outputs[0]) as im:
        assert im.format == "JPEG" and im.size == (1190, 1684)
    with pytest.raises(InvalidInputError):
        pdf_to_images(src, tmp_path / "x", pages_spec="9")
    with pytest.raises(InvalidInputError):
        pdf_to_images(src, tmp_path / "x", dpi=5)


def test_pdf_to_images_progress(tmp_path, pdf_factory):
    src = pdf_factory("doc.pdf", pages=3)
    seen = []
    pdf_to_images(src, tmp_path / "o", dpi=50, ctx=JobContext(progress=lambda d, t, m: seen.append((d, t))))
    assert seen == [(0, 3), (1, 3), (2, 3)]


# --------------------------------------------------------- PDF -> image -> PDF
def test_pdf_image_pdf_rasterises_and_keeps_geometry(tmp_path, pdf_factory):
    src = pdf_factory("doc.pdf", pages=3, size=(400, 300))
    out = tmp_path / "r.pdf"
    res = pdf_to_image_to_pdf(src, out, dpi=100)
    with pymupdf.open(out) as d:
        assert d.page_count == 3
        assert all(round(p.rect.width) == 400 and round(p.rect.height) == 300 for p in d)
        assert page_texts(out) == ["", "", ""]          # text is gone: pages are pictures
        assert all(p.get_images() for p in d)
    assert any("Selectable" in w for w in res.warnings)


# --------------------------------------------------------------------------- HTML
def test_html_to_pdf_text_table_image_and_multipage(tmp_path):
    Image.new("RGB", (60, 40), "green").save(tmp_path / "pic.png")
    rows = "".join(f"<tr><td>Row {i}</td><td>{i * 3}</td></tr>" for i in range(150))
    html = (f"<html><body><h1>Quarterly Report</h1><img src='pic.png' width='60'>"
            f"<table border='1'>{rows}</table></body></html>")
    f = tmp_path / "r.html"
    f.write_text(html, encoding="utf-8")
    out = tmp_path / "r.pdf"
    res = html_to_pdf(out, html_file=f)
    assert res.details["pages"] >= 2
    with pymupdf.open(out) as d:
        assert "Quarterly Report" in d[0].get_text()
        assert d[0].get_images()                          # relative image resolved locally
        assert "Row 149" in d[d.page_count - 1].get_text()
    out2 = tmp_path / "land.pdf"
    html_to_pdf(out2, html_text="<p>Hi</p>", orientation="landscape", page_size="letter")
    with pymupdf.open(out2) as d:
        assert (round(d[0].rect.width), round(d[0].rect.height)) == (792, 612)


def test_html_warnings_and_errors(tmp_path):
    w = scan_html("<script>1</script><img src='https://x.test/a.png'><link href='../s.css'>", tmp_path)
    assert len(w) == 3
    with pytest.raises(InvalidInputError):
        html_to_pdf(tmp_path / "x.pdf")
    with pytest.raises(InvalidInputError):
        html_to_pdf(tmp_path / "x.pdf", html_text="   ")
    with pytest.raises(InvalidInputError):
        html_to_pdf(tmp_path / "x.pdf", html_text="<p>x</p>", margin_mm=300)
    bad = tmp_path / "x.txt"
    bad.write_text("x")
    with pytest.raises(InvalidInputError):
        html_to_pdf(tmp_path / "x.pdf", html_file=bad)


# ------------------------------------------------------------ Office engine logic
def test_missing_engine_gives_install_help(monkeypatch, tmp_path):
    from core import office_engine
    monkeypatch.setattr(office_engine, "find_soffice", lambda: None)
    with pytest.raises(EngineMissingError, match="libreoffice.org"):
        office_engine.require_soffice()
    docx_like = tmp_path / "x.docx"
    import zipfile
    with zipfile.ZipFile(docx_like, "w") as z:
        z.writestr("word/document.xml", "<w/>")
    with pytest.raises(EngineMissingError):  # the public converter surfaces the same helpful error
        word_to_pdf(docx_like, tmp_path / "o.pdf")


def test_office_validation_rejects_fakes(tmp_path):
    fake = tmp_path / "fake.docx"
    fake.write_text("this is plain text, not a zip")
    with pytest.raises(InvalidInputError, match="damaged|not a real"):
        word_to_pdf(fake, tmp_path / "o.pdf")
    txt = tmp_path / "a.txt"
    txt.write_text("x")
    with pytest.raises(InvalidInputError, match="not supported"):
        word_to_pdf(txt, tmp_path / "o.pdf")
    macro = tmp_path / "m.xlsm"
    macro.write_bytes(b"PK")
    with pytest.raises(InvalidInputError, match="not supported"):
        excel_to_pdf(macro, tmp_path / "o.pdf")


# ----------------------------------------------------- Office conversions (LibreOffice)
@pytest.fixture
def docx_file(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Annual Summary", 1)
    d.add_paragraph("First page paragraph.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 1).text = "Alpha", "Omega"
    d.add_page_break()
    d.add_paragraph("Second page text.")
    p = tmp_path / "doc.docx"
    d.save(p)
    return p


@needs_lo
def test_word_to_pdf_real_render(tmp_path, docx_file):
    out = tmp_path / "w.pdf"
    res = word_to_pdf(docx_file, out)
    texts = page_texts(out)
    assert len(texts) == 2                                  # explicit page break kept
    assert "Annual Summary" in texts[0] and "Alpha" in texts[0] and "Omega" in texts[0]
    assert "Second page text." in texts[1]
    assert any("Fonts" in w or "fonts" in w for w in res.warnings)


@needs_lo
def test_powerpoint_to_pdf_one_page_per_slide(tmp_path):
    import pptx
    prs = pptx.Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Slide number {i + 1}"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    out = tmp_path / "p.pdf"
    powerpoint_to_pdf(p, out)
    texts = page_texts(out)
    assert len(texts) == 3 and "Slide number 2" in texts[1]
    with pymupdf.open(out) as d:  # slide dimensions (default 10in x 7.5in) are preserved
        assert abs(d[0].rect.width - 720) < 2 and abs(d[0].rect.height - 540) < 2


@pytest.fixture
def xlsx_file(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Total"])
    ws.append(["North", "=1+2"])
    ws2 = wb.create_sheet("Costs")
    ws2.append(["Item", "Cost"])
    ws2.append(["Rent", 500])
    secret = wb.create_sheet("Secret")
    secret.append(["TOPSECRET"])
    secret.sheet_state = "hidden"
    p = tmp_path / "book.xlsx"
    wb.save(p)
    return p


def test_list_sheets(xlsx_file):
    sheets = list_sheets(xlsx_file)
    assert [(s.name, s.hidden) for s in sheets] == [("Sales", False), ("Costs", False), ("Secret", True)]


@needs_lo
def test_excel_to_pdf_all_and_selected_sheets(tmp_path, xlsx_file):
    out = tmp_path / "x.pdf"
    res = excel_to_pdf(xlsx_file, out)
    text = " ".join(page_texts(out))
    assert "North" in text and "Rent" in text and "TOPSECRET" not in text  # hidden sheet is not exported
    assert "3" in text                                                        # formula evaluated by the engine
    assert any("Hidden sheet" in w for w in res.warnings)
    out2 = tmp_path / "x2.pdf"
    res2 = excel_to_pdf(xlsx_file, out2, sheets=["Costs"], orientation="landscape", fit_to_width=True)
    text2 = " ".join(page_texts(out2))
    assert "Rent" in text2 and "North" not in text2
    assert any("re-saved" in w for w in res2.warnings)
    with pymupdf.open(out2) as d:
        assert d[0].rect.width > d[0].rect.height
    with pytest.raises(InvalidInputError, match="Sheet not found"):
        excel_to_pdf(xlsx_file, tmp_path / "x3.pdf", sheets=["Nope"])


@needs_lo
def test_office_cancel_kills_engine(tmp_path, docx_file):
    ctx = JobContext()
    ctx.cancel.cancel()
    with pytest.raises(OperationCancelled):
        word_to_pdf(docx_file, tmp_path / "c.pdf", ctx=ctx)
    assert not (tmp_path / "c.pdf").exists()


@needs_lo
def test_office_corrupt_zip_reports_conversion_problem(tmp_path):
    import zipfile
    bad = tmp_path / "bad.docx"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("word/document.xml", "<broken")  # right marker, garbage content
    with pytest.raises(ToolkitError) as ei:   # a friendly error, never a raw exception
        word_to_pdf(bad, tmp_path / "b.pdf")
    assert "Traceback" not in str(ei.value) and not (tmp_path / "b.pdf").exists()
