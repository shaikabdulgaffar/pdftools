"""Drives the real Qt workspaces headlessly (QT_QPA_PLATFORM=offscreen): every tool, through the widgets,
on a worker thread, verified by inspecting the files it wrote."""
import os
import time
import zipfile
from pathlib import Path

import pymupdf
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QCheckBox, QComboBox, QLineEdit, QListWidget, QSpinBox  # noqa: E402

from app.settings import AppSettings  # noqa: E402
from app.theme import stylesheet  # noqa: E402
from app.tools import TOOLS, TOOLS_BY_ID, describe_file  # noqa: E402
from core.office_engine import find_soffice  # noqa: E402
from ui import tool_workspace  # noqa: E402
from ui.tool_workspace import ToolWorkspace  # noqa: E402
from conftest import make_pdf, page_texts  # noqa: E402

needs_lo = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(stylesheet("light"))
    yield app
    # Tear every widget (and any worker threads) down while the QApplication still exists; otherwise Python's
    # interpreter-exit garbage collection destroys them afterwards and Qt aborts.
    for w in app.topLevelWidgets():
        w.close()
        w.deleteLater()
    for _ in range(5):
        app.processEvents()
        from PySide6.QtCore import QCoreApplication, QEvent
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


@pytest.fixture
def settings(tmp_path, monkeypatch):
    from PySide6.QtCore import QSettings
    QSettings.setPath(QSettings.NativeFormat, QSettings.UserScope, str(tmp_path / "qs"))
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path / "qs"))
    return AppSettings()


def wait(app, cond, timeout=90):
    end = time.time() + timeout
    while time.time() < end:
        app.processEvents()
        if cond():
            return
        time.sleep(0.01)
    raise AssertionError("timed out waiting for the GUI")


def idle(ws):
    return ws._thread is None and ws._rerun is None


def set_opt(ws, key, value):
    w = ws._opt_widgets[key]
    if isinstance(w, QComboBox):
        w.setCurrentIndex(w.findData(value))
        assert w.currentData() == value
    elif isinstance(w, QSpinBox):
        w.setValue(value)
    elif isinstance(w, QCheckBox):
        w.setChecked(value)
    elif isinstance(w, QLineEdit):
        w.setText(value)


def run(app, ws, out: Path | None = None, expect_ok=True):
    if out is not None:
        ws.output_edit.setText(str(out))
    assert ws.btn_run.isEnabled()
    ws.btn_run.click()
    wait(app, lambda: idle(ws))
    ok = ws.result_title.text().startswith("✓")
    assert ok == expect_ok, ws.result_title.text() + " | " + ws.result_msg.text()
    return ws


def make(app, settings, tool_id):
    ws = ToolWorkspace(TOOLS_BY_ID[tool_id], settings)
    ws.show()
    return ws


def test_registry_has_all_16_tools_with_distinct_workspaces(qapp, settings):
    assert len(TOOLS) == 16
    cats = {}
    for t in TOOLS:
        cats.setdefault(t.category, []).append(t.title)
    assert cats == {
        "PDF Management": ["Merge PDF", "Split PDF", "Remove Pages", "Extract Pages", "Organize PDF", "Compress PDF"],
        "Convert to PDF": ["Image to PDF", "Word to PDF", "PowerPoint to PDF", "Excel to PDF", "HTML to PDF"],
        "Convert from PDF": ["PDF to Image", "PDF to Image to PDF"],
        "Security": ["Protect PDF", "Unlock PDF"],
        "Image Tools": ["Image Compressor"]}
    for t in TOOLS:
        ws = make(qapp, settings, t.id)
        assert ws.btn_run.text() == t.title and not ws.btn_run.isEnabled()  # disabled until input is given


def test_main_window_navigation(qapp, settings):
    from app.main_window import MainWindow
    w = MainWindow(settings, "x.log")
    w.show()
    for t in TOOLS:
        w.open_tool(t.id)
        assert w.stack.currentWidget() is w._workspaces[t.id]
    for name in ("Home", "PDF Management", "Convert to PDF", "Convert from PDF", "Security", "Image Tools",
                 "Settings", "About"):
        w.show_page(name)
        assert w.stack.currentWidget() is w._pages[name]
    w.apply_theme("dark")
    w.apply_theme("light")


def test_merge_via_gui_reorder_and_overwrite_confirmation(qapp, settings, tmp_path, monkeypatch):
    a, b, c = (make_pdf(tmp_path / f"{n}.pdf", 1, n.upper()) for n in "abc")
    ws = make(qapp, settings, "merge")
    ws._add_files([a, b, c])
    ws.list.item(2).setSelected(True)
    ws.btn_up.click()
    ws.btn_up.click()  # c moves to the top
    assert [Path(ws.list.item(i).data(Qt.UserRole)).name for i in range(3)] == ["c.pdf", "a.pdf", "b.pdf"]
    assert "3 page" not in ws.list.item(0).text() and "1 page(s)" in ws.list.item(0).text()
    out = tmp_path / "result.pdf"
    run(qapp, ws, out)
    assert page_texts(out) == ["C 1", "A 1", "B 1"]
    assert not any(p.name.startswith(".") for p in tmp_path.iterdir())
    # second run: existing file -> user is asked; declining keeps it, accepting replaces it
    asked = []
    monkeypatch.setattr(tool_workspace, "confirm_overwrite", lambda parent, paths: asked.append(paths) or False)
    ws.btn_run.click()
    wait(qapp, lambda: idle(ws))
    assert asked and "nothing was replaced" in ws.status.text().lower()
    ws.list.takeItem(2)
    monkeypatch.setattr(tool_workspace, "confirm_overwrite", lambda parent, paths: True)
    ws.btn_run.click()
    wait(qapp, lambda: idle(ws))
    assert ws.result_title.text().startswith("✓") and len(page_texts(out)) == 2


def test_merge_needs_two_files_message(qapp, settings, tmp_path):
    ws = make(qapp, settings, "merge")
    ws._add_files([make_pdf(tmp_path / "a.pdf")])
    ws.output_edit.setText(str(tmp_path / "o.pdf"))
    run(qapp, ws, expect_ok=False)
    assert "at least 2" in ws.result_msg.text()
    assert not (tmp_path / "o.pdf").exists()


def test_unsupported_and_corrupt_inputs_are_reported_not_crashing(qapp, settings, tmp_path):
    ws = make(qapp, settings, "merge")
    txt = tmp_path / "x.txt"
    txt.write_text("hi")
    ws._add_files([txt])
    assert ws.list.count() == 0 and "not supported" in ws.status.text()
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.4 junk")
    ws._add_files([bad, make_pdf(tmp_path / "ok.pdf")])
    run(qapp, ws, tmp_path / "o.pdf", expect_ok=False)
    assert "damaged" in ws.result_msg.text()
    assert ws.btn_details.isVisible() and "InvalidInputError" in ws._error_details
    assert ws.btn_run.isEnabled()  # app still usable after failure


def test_split_via_gui_ranges(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 10)
    ws = make(qapp, settings, "split")
    ws._add_files([src])
    assert ws.output_edit.text().endswith("doc_split")
    set_opt(ws, "mode", "ranges")
    assert ws.form.isRowVisible(ws._opt_widgets["ranges"]) and not ws.form.isRowVisible(ws._opt_widgets["every_n"])
    set_opt(ws, "ranges", "1-3, 5, 8-10")
    run(qapp, ws, tmp_path / "parts")
    assert sorted(p.name for p in (tmp_path / "parts").iterdir()) == [
        "doc_page_005.pdf", "doc_pages_001-003.pdf", "doc_pages_008-010.pdf"]
    assert ws.btn_open.isVisible()


def test_remove_and_extract_with_thumbnails_and_selection(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 5)
    ws = make(qapp, settings, "remove")
    ws._add_files([src])
    wait(qapp, lambda: ws.grid.count() == 5 and len(ws.grid._base) == 5)  # thumbnails rendered off-thread
    ws.grid.item(1).setSelected(True)
    ws.grid.item(3).setSelected(True)
    assert ws._opt_widgets["pages"].text() == "2, 4"  # selection fills the page field
    run(qapp, ws, tmp_path / "removed.pdf")
    assert page_texts(tmp_path / "removed.pdf") == ["Page 1", "Page 3", "Page 5"]
    ws2 = make(qapp, settings, "extract")
    ws2._add_files([src])
    wait(qapp, lambda: ws2.grid.count() == 5)
    set_opt(ws2, "pages", "4, 1-2")
    run(qapp, ws2, tmp_path / "ext.pdf")
    assert page_texts(tmp_path / "ext.pdf") == ["Page 1", "Page 2", "Page 4"]
    set_opt(ws2, "pages", "99")
    run(qapp, ws2, tmp_path / "ext2.pdf", expect_ok=False)
    assert "Page 99 does not exist" in ws2.result_msg.text()


def test_organize_via_gui_rotate_duplicate_delete_move_undo_reset(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 4)
    ws = make(qapp, settings, "organize")
    ws._add_files([src])
    g = ws.grid
    wait(qapp, lambda: g.count() == 4 and len(g._base) == 4)
    g.item(0).setSelected(True)
    g.rotate_selected(90)
    g.duplicate_selected()                      # [0r, 0r, 1, 2, 3]
    assert [s for s, _ in g._snapshot()] == [0, 0, 1, 2, 3]
    g.clearSelection()
    g.item(4).setSelected(True)
    g.move_selected(-1)                         # last page one place earlier
    assert [s for s, _ in g._snapshot()] == [0, 0, 1, 3, 2]
    g.clearSelection()
    g.item(2).setSelected(True)
    g.delete_selected()                         # drop original page 2
    assert [s for s, _ in g._snapshot()] == [0, 0, 3, 2]
    g.undo()
    assert g.count() == 5
    g.reset_plan()
    assert [s for s, _ in g._snapshot()] == [0, 1, 2, 3] and all(r == 0 for _, r in g._snapshot())
    g.undo()  # undo the reset
    g.clearSelection()
    g.item(2).setSelected(True)
    g.delete_selected()
    run(qapp, ws, tmp_path / "org.pdf")
    assert page_texts(tmp_path / "org.pdf") == ["Page 1", "Page 1", "Page 4", "Page 3"]
    with pymupdf.open(tmp_path / "org.pdf") as d:
        assert [p.rotation for p in d] == [90, 90, 0, 0]
    assert len(page_texts(src)) == 4  # original untouched


def test_compress_pdf_via_gui(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 3)
    ws = make(qapp, settings, "compress")
    ws._add_files([src])
    set_opt(ws, "preset", "high")
    run(qapp, ws, tmp_path / "c.pdf")
    assert "→" in ws.result_msg.text()  # before/after sizes reported
    assert not ws.btn_cancel.isEnabled() or not ws.btn_cancel.isVisible()


def test_image_to_pdf_via_gui(qapp, settings, img_dir, tmp_path):
    ws = make(qapp, settings, "img2pdf")
    ws._add_files([img_dir / "photo.jpg", img_dir / "alpha.png", img_dir / "multi.tif"])
    assert "800×600" in ws.list.item(0).text()
    set_opt(ws, "page_size", "letter")
    assert ws.form.isRowVisible(ws._opt_widgets["orientation"])      # visible for fixed page sizes
    set_opt(ws, "page_size", "image")
    assert not ws.form.isRowVisible(ws._opt_widgets["orientation"])
    run(qapp, ws, tmp_path / "i.pdf")
    with pymupdf.open(tmp_path / "i.pdf") as d:
        assert d.page_count == 4


def test_html_to_pdf_via_gui_both_modes(qapp, settings, tmp_path):
    ws = make(qapp, settings, "html2pdf")
    f = tmp_path / "page.html"
    f.write_text("<h1>From File</h1><script>1</script>", encoding="utf-8")
    ws._add_files([f])
    run(qapp, ws, tmp_path / "f.pdf")
    assert "From File" in page_texts(tmp_path / "f.pdf")[0] and "Scripts are not executed" in ws.result_warn.text()
    ws.html_mode_text.setChecked(True)
    ws.html_edit.setPlainText("<h2>Typed HTML</h2>")
    run(qapp, ws, tmp_path / "t.pdf")
    assert "Typed HTML" in page_texts(tmp_path / "t.pdf")[0]


def test_pdf_to_image_and_rasterize_via_gui(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 3)
    ws = make(qapp, settings, "pdf2img")
    ws._add_files([src])
    set_opt(ws, "format", "jpeg")
    set_opt(ws, "dpi", 72)
    set_opt(ws, "pages", "2-3")
    run(qapp, ws, tmp_path / "imgs")
    assert sorted(p.name for p in (tmp_path / "imgs").iterdir()) == ["doc_page_002.jpg", "doc_page_003.jpg"]
    ws2 = make(qapp, settings, "rasterize")
    ws2._add_files([src])
    run(qapp, ws2, tmp_path / "flat.pdf")
    assert page_texts(tmp_path / "flat.pdf") == ["", "", ""] and "Selectable" in ws2.result_warn.text()


def test_protect_and_unlock_via_gui(qapp, settings, tmp_path):
    src = make_pdf(tmp_path / "doc.pdf", 2)
    ws = make(qapp, settings, "protect")
    ws._add_files([src])
    set_opt(ws, "password", "hunter2")
    set_opt(ws, "confirm", "different")
    run(qapp, ws, tmp_path / "p.pdf", expect_ok=False)
    assert "do not match" in ws.result_msg.text() and not (tmp_path / "p.pdf").exists()
    set_opt(ws, "confirm", "hunter2")
    set_opt(ws, "owner", "own-pw")
    set_opt(ws, "allow_copy", False)
    run(qapp, ws, tmp_path / "p.pdf")
    assert ws._opt_widgets["password"].text() == ""  # secrets cleared after use
    with pymupdf.open(tmp_path / "p.pdf") as d:
        assert d.needs_pass and d.authenticate("hunter2") == 2 and not d.permissions & pymupdf.PDF_PERM_COPY
    un = make(qapp, settings, "unlock")
    un._add_files([tmp_path / "p.pdf"])
    set_opt(un, "password", "wrong")
    run(qapp, un, tmp_path / "u.pdf", expect_ok=False)
    assert "incorrect" in un.result_msg.text() and not (tmp_path / "u.pdf").exists()
    set_opt(un, "password", "hunter2")
    run(qapp, un, tmp_path / "u.pdf")
    assert page_texts(tmp_path / "u.pdf") == ["Page 1", "Page 2"]
    # a management tool refuses the protected file and points to Unlock
    mg = make(qapp, settings, "extract")
    mg._add_files([tmp_path / "p.pdf"])
    set_opt(mg, "pages", "1")
    run(qapp, mg, tmp_path / "x.pdf", expect_ok=False)
    assert "Unlock PDF" in mg.result_msg.text()


def test_image_compressor_via_gui_with_table_and_overwrite_confirmation(qapp, settings, img_dir, tmp_path, monkeypatch):
    import shutil
    work = tmp_path / "w"
    shutil.copytree(img_dir, work)
    ws = make(qapp, settings, "imgcompress")
    ws._add_files([work / "photo.jpg", work / "alpha.png", work / "broken.png"])
    set_opt(ws, "quality", 40)
    set_opt(ws, "max_width", 400)
    before = (work / "photo.jpg").read_bytes()
    run(qapp, ws, tmp_path / "out")
    assert ws.result_table.isVisible() and ws.result_table.rowCount() == 2
    assert "broken.png" in ws.result_warn.text()
    assert (work / "photo.jpg").read_bytes() == before and (tmp_path / "out" / "photo_compressed.jpg").exists()
    # overwrite mode asks first; declining changes nothing
    set_opt(ws, "overwrite", True)
    assert not ws.output_edit.isEnabled()
    monkeypatch.setattr(tool_workspace, "confirm", lambda *a, **k: False)
    ws.btn_run.click()
    wait(qapp, lambda: idle(ws))
    assert (work / "photo.jpg").read_bytes() == before
    monkeypatch.setattr(tool_workspace, "confirm", lambda *a, **k: True)
    run(qapp, ws)
    assert (work / "photo.jpg").read_bytes() != before


@needs_lo
def test_office_tools_via_gui(qapp, settings, tmp_path):
    import docx, openpyxl, pptx
    d = docx.Document(); d.add_paragraph("Hello from Word"); d.save(tmp_path / "a.docx")
    p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Deck"; p.save(tmp_path / "a.pptx")
    wb = openpyxl.Workbook(); wb.active.title = "One"; wb.active["A1"] = "alpha"; wb.create_sheet("Two")["A1"] = "beta"
    wb.save(tmp_path / "a.xlsx")
    for tid, f, needle in [("word2pdf", "a.docx", "Hello from Word"), ("ppt2pdf", "a.pptx", "Deck")]:
        ws = make(qapp, settings, tid)
        ws._add_files([tmp_path / f])
        run(qapp, ws, tmp_path / f"{tid}.pdf")
        assert needle in page_texts(tmp_path / f"{tid}.pdf")[0]
    ws = make(qapp, settings, "xls2pdf")
    ws._add_files([tmp_path / "a.xlsx"])
    sheets: QListWidget = ws._opt_widgets["sheets"]
    assert [sheets.item(i).text() for i in range(sheets.count())] == ["One", "Two"]
    sheets.item(0).setCheckState(Qt.Unchecked)  # export only "Two"
    run(qapp, ws, tmp_path / "x.pdf")
    t = " ".join(page_texts(tmp_path / "x.pdf"))
    assert "beta" in t and "alpha" not in t


def test_describe_file_handles_everything(tmp_path, img_dir):
    pdf = make_pdf(tmp_path / "d.pdf", 2)
    assert describe_file(pdf).startswith("2 page(s)")
    assert "px" in describe_file(img_dir / "photo.jpg")
    assert "Cannot read" in describe_file(img_dir / "broken.png")
